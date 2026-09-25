#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ETL COMPLEMENTARIO — Sección 1 (Transformaciones PND) derivada de BASE DETALLE MENSUAL.

Contexto
--------
El archivo canónico de la Sec 1 (`Inversiones 2026 - PND 2022-2026.xlsx`, hoja `Base`)
trae por proyecto las columnas `Transformación` y `Componente`. Cuando ese archivo NO
llega en la entrega de un corte, la información monetaria equivalente sí está en
`BASE DETALLE MENSUAL INVERSIÓN 2018-2026.xlsx` (hoja `BASE`), filtrando `Año` + `Mes`
del corte: se verificó que el universo coincide EXACTO en total y por sector
(1.366 filas, 88.401,2 mmm vigente para 2026/MAR).

Lo único que `BASE` NO trae es la dimensión Transformación/Componente (ni `BPIN`).
Este ETL la reconstruye de forma **APROXIMADA** cruzando por la llave programática
(`Código, Gsto, Prog, Subp, Proy, Rec, Sit`) contra un archivo catálogo (el
`Inversiones` de un corte previo, típicamente marzo). En la validación sobre marzo el
cruce empareja el 100% del vigente y reproduce el desglose por transformación con
error < 1% por transformación (8 llaves programáticas son ambiguas porque solo se
distinguen por `BPIN`/`SubP2`/`SubOrd`, ausentes en `BASE`).

⚠️ Es una APROXIMACIÓN. El resultado exacto requiere el archivo `Inversiones` del corte
   o que `BASE` incluya la columna `BPIN`. No sustituye a `load_bitacora_excel.py`.

No modifica ningún ETL existente ni el esquema.

Uso
---
    python etl/load_sec1_desde_base.py \
        --base    "C:\\...\\Junio\\6. EJECUCIÓN SECTORIAL\\BASE DETALLE MENSUAL INVERSIÓN 2018-2026.xlsx" \
        --crosswalk "C:\\...\\Marzo\\1. INVERSIONES 2026 - PND 2022 - 2026\\Inversiones 2026 - PND 2022-2026.xlsx" \
        --anio 2026 --mes JUN \
        --crear-bitacora --numero 3 --periodo 2026-II --corte 2026-06-30
"""
import argparse
import sqlite3
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

# La consola de Windows suele ser cp1252 y no puede imprimir → ni ⚠; forzamos UTF-8.
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

MMM = 1_000_000_000
TOP_COMPONENTES = 5

DEFAULT_DB = str(Path(__file__).resolve().parent.parent / "db" / "pgn.db")

# ── Índices de columna en la hoja BASE (encabezado en fila 3; datos desde fila 4) ──
B_ANIO, B_MES = 0, 1
B_COD, B_GSTO, B_PROG, B_SUBP, B_PROY = 4, 8, 9, 10, 11
B_REC, B_SIT = 13, 14
B_VIG, B_COM, B_OBL, B_PAG = 19, 21, 22, 23


def s_(v):
    return "" if v is None else str(v).strip()


def up(v):
    return unicodedata.normalize("NFC", s_(v)).upper()


def to_float(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def llave_prog(cod, gsto, prog, subp, proy, rec, sit):
    return (s_(cod), s_(gsto), s_(prog), s_(subp), s_(proy), s_(rec), s_(sit))


# ── 1) Crosswalk llave programática → (transformación, componente) ────────────
def build_crosswalk(inv_path):
    """Lee la hoja `Base` de Inversiones y devuelve:
        xwalk: dict  llave -> (transformacion, componente)   (desambiguada por mayor vigente)
        n_amb: número de llaves con >1 transformación distinta
    """
    wb = openpyxl.load_workbook(inv_path, data_only=True, read_only=True)
    ws = wb["Base"]
    hdr = [s_(h) for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    idx = {h: i for i, h in enumerate(hdr)}

    def gi(*names):
        for n in names:
            if n in idx:
                return idx[n]
        raise KeyError(f"No se encontró ninguna de las columnas {names} en hoja Base")

    iCod = gi("Código"); iGsto = gi("Gsto"); iProg = gi("Prog"); iSubp = gi("Subp")
    iProy = gi("Proy"); iRec = gi("Rec"); iSit = gi("Sit")
    iVig = gi("Vigente"); iT = gi("Transformación", "Transformacion")
    iC = gi("Componente")

    # llave -> {(t,c): vigente acumulado}
    acc = defaultdict(lambda: defaultdict(float))
    tset = defaultdict(set)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[iCod] is None and not s_(row[iT]):
            continue
        t = s_(row[iT])
        if not t:
            continue
        c = s_(row[iC])
        k = llave_prog(row[iCod], row[iGsto], row[iProg], row[iSubp],
                       row[iProy], row[iRec], row[iSit])
        acc[k][(t, c)] += to_float(row[iVig])
        tset[k].add(t)
    wb.close()

    xwalk = {}
    for k, opts in acc.items():
        # elige la (transformación, componente) con mayor vigente para esa llave
        xwalk[k] = max(opts.items(), key=lambda kv: kv[1])[0]
    n_amb = sum(1 for k, ts in tset.items() if len(ts) > 1)
    return xwalk, n_amb


# ── 2) Agrega BASE (Año+Mes) asignando transformación/componente por la llave ──
def aggregate_base(base_path, anio, mes, xwalk):
    mes = mes.upper()
    wb = openpyxl.load_workbook(base_path, data_only=True, read_only=True)
    ws = wb["BASE"]

    t_acc = defaultdict(lambda: {"v": 0.0, "c": 0.0, "o": 0.0, "p": 0.0})
    tc_acc = defaultdict(lambda: defaultdict(float))
    matched = unmatched = 0.0
    n_rows = n_unmatched_rows = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[B_ANIO] != anio or up(row[B_MES]) != mes:
            continue
        n_rows += 1
        v = to_float(row[B_VIG]); c = to_float(row[B_COM])
        o = to_float(row[B_OBL]); p = to_float(row[B_PAG])
        k = llave_prog(row[B_COD], row[B_GSTO], row[B_PROG], row[B_SUBP],
                       row[B_PROY], row[B_REC], row[B_SIT])
        tc = xwalk.get(k)
        if tc is None:
            unmatched += v
            n_unmatched_rows += 1
            continue
        t, comp = tc
        matched += v
        t_acc[t]["v"] += v; t_acc[t]["c"] += c
        t_acc[t]["o"] += o; t_acc[t]["p"] += p
        if comp:
            tc_acc[t][comp] += v
    wb.close()
    return t_acc, tc_acc, matched, unmatched, n_rows, n_unmatched_rows


# ── 3) Carga las 3 tablas de la Sec 1 (misma lógica que load_bitacora_excel) ──
def load_sec1(conn, bid, vigencia, t_acc, tc_acc):
    conn.execute("DELETE FROM inversion_transformaciones WHERE bitacora_id=? AND vigencia=?", (bid, vigencia))
    conn.execute("DELETE FROM inversion_componentes_pnd  WHERE bitacora_id=? AND vigencia=?", (bid, vigencia))
    conn.execute("DELETE FROM ejecucion_transformaciones WHERE bitacora_id=? AND vigencia=?", (bid, vigencia))

    total_v = sum(d["v"] for d in t_acc.values()) or 1

    rows = [
        (bid, vigencia, t, round(d["v"] / MMM, 3), round(d["v"] / total_v * 100, 2))
        for t, d in t_acc.items()
    ]
    conn.executemany(
        """INSERT OR REPLACE INTO inversion_transformaciones
               (bitacora_id, vigencia, transformador, inversion_mmm, peso_pct)
           VALUES (?,?,?,?,?)""", rows)
    print(f"  [OK] inversion_transformaciones : {len(rows):>4} filas")

    rows_c = []
    for t, comps in tc_acc.items():
        t_total = t_acc[t]["v"] or 1
        sorted_c = sorted(comps.items(), key=lambda x: -x[1])
        top = sorted_c[:TOP_COMPONENTES]
        otros_v = sum(v for _, v in sorted_c[TOP_COMPONENTES:])
        for comp, v in top:
            rows_c.append((bid, vigencia, t, comp, round(v / MMM, 3), round(v / t_total * 100, 2)))
        if otros_v > 0:
            rows_c.append((bid, vigencia, t, "OTROS COMPONENTES", round(otros_v / MMM, 3),
                           round(otros_v / t_total * 100, 2)))
    conn.executemany(
        """INSERT OR REPLACE INTO inversion_componentes_pnd
               (bitacora_id, vigencia, transformador, componente, vigente_mmm, peso_pct)
           VALUES (?,?,?,?,?,?)""", rows_c)
    print(f"  [OK] inversion_componentes_pnd  : {len(rows_c):>4} filas")

    rows_e = []
    for t, d in t_acc.items():
        v = d["v"] or 1
        rows_e.append((
            bid, vigencia, t,
            round(d["v"] / MMM, 3), round(d["c"] / MMM, 3),
            round(d["o"] / MMM, 3), round(d["p"] / MMM, 3),
            round(d["c"] / v * 100, 1), round(d["o"] / v * 100, 1), round(d["p"] / v * 100, 1),
        ))
    conn.executemany(
        """INSERT OR REPLACE INTO ejecucion_transformaciones
               (bitacora_id, vigencia, transformador,
                apr_vigente_mmm, compromisos_mmm, obligaciones_mmm, pagos_mmm,
                pct_c_av, pct_o_av, pct_p_av)
           VALUES (?,?,?,?,?,?,?,?,?,?)""", rows_e)
    print(f"  [OK] ejecucion_transformaciones : {len(rows_e):>4} filas")


def resolver_bitacora(conn, args):
    if args.bitacora_id is not None:
        return args.bitacora_id
    if args.crear_bitacora:
        row = conn.execute("SELECT id FROM metadatos_bitacora WHERE periodo=?", (args.periodo,)).fetchone()
        if row:
            print(f"Bitácora {args.periodo} ya existe (id={row[0]}); se reutiliza.")
            return row[0]
        cur = conn.execute(
            """INSERT INTO metadatos_bitacora
                   (numero_bitacora, periodo, corte_fecha, fuente_principal, notas)
               VALUES (?,?,?,?,?)""",
            (args.numero, args.periodo, args.corte,
             "BASE DETALLE MENSUAL (Sec 1 aproximada por llave programática)",
             args.notas or "Sec 1 derivada de BASE — desglose por transformación aproximado."))
        conn.commit()
        print(f"Bitácora creada: {args.periodo} (id={cur.lastrowid}).")
        return cur.lastrowid
    row = conn.execute("SELECT id FROM metadatos_bitacora ORDER BY corte_fecha DESC LIMIT 1").fetchone()
    if not row:
        raise SystemExit("No hay bitácoras. Usa --crear-bitacora --numero N --periodo P --corte YYYY-MM-DD.")
    return row[0]


def main():
    ap = argparse.ArgumentParser(description="Carga Sec 1 (Transformaciones PND) aproximada desde BASE DETALLE MENSUAL.")
    ap.add_argument("--base", required=True, help="Ruta de BASE DETALLE MENSUAL INVERSIÓN ....xlsx")
    ap.add_argument("--crosswalk", required=True, help="Archivo Inversiones (catálogo llave→transformación/componente).")
    ap.add_argument("--anio", type=int, default=2026)
    ap.add_argument("--mes", default="JUN", help="Mes de corte (ENE..DIC).")
    ap.add_argument("--db", default=DEFAULT_DB)
    ap.add_argument("--bitacora-id", type=int, default=None)
    ap.add_argument("--crear-bitacora", action="store_true")
    ap.add_argument("--numero", type=int, default=None)
    ap.add_argument("--periodo", default=None)
    ap.add_argument("--corte", default=None)
    ap.add_argument("--notas", default=None)
    args = ap.parse_args()

    if args.crear_bitacora and not (args.numero and args.periodo and args.corte):
        raise SystemExit("--crear-bitacora requiere --numero, --periodo y --corte.")

    print("Construyendo crosswalk llave programática → transformación/componente ...")
    xwalk, n_amb = build_crosswalk(args.crosswalk)
    print(f"  {len(xwalk)} llaves en el catálogo; {n_amb} ambiguas (desambiguadas por mayor vigente).")

    print(f"Agregando BASE filtrada a Año={args.anio}, Mes={args.mes.upper()} ...")
    t_acc, tc_acc, matched, unmatched, n_rows, n_unm = aggregate_base(args.base, args.anio, args.mes, xwalk)
    tot = matched + unmatched or 1
    print(f"  Filas del corte: {n_rows}  |  emparejadas: {n_rows - n_unm}  |  sin llave: {n_unm}")
    print(f"  Vigente emparejado : {matched/MMM:>12,.1f} mmm ({matched/tot*100:5.1f}%)")
    print(f"  Vigente SIN mapear : {unmatched/MMM:>12,.1f} mmm ({unmatched/tot*100:5.1f}%)  ← no entra a Sec 1")

    conn = sqlite3.connect(args.db)
    conn.execute("PRAGMA foreign_keys = ON")
    bid = resolver_bitacora(conn, args)
    print(f"Bitácora destino: id={bid}, vigencia={args.anio}")
    load_sec1(conn, bid, args.anio, t_acc, tc_acc)
    conn.commit()

    print("\n--- Desglose por transformación (mmm vigente) ---")
    for r in conn.execute(
        """SELECT transformador, inversion_mmm, peso_pct
             FROM inversion_transformaciones
            WHERE bitacora_id=? AND vigencia=? ORDER BY inversion_mmm DESC""", (bid, args.anio)):
        print(f"  {r[0][:52]:<54}{r[1]:>10,.1f}  {r[2]:>5.1f}%")
    conn.close()
    print("\n⚠️  Desglose por transformación APROXIMADO (cruce por llave programática). "
          "Para exactitud usar el archivo Inversiones del corte o agregar BPIN a BASE.")


if __name__ == "__main__":
    main()
